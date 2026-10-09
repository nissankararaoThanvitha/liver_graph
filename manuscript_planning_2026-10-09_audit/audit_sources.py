from pathlib import Path
import hashlib,json,csv,re,collections
import pandas as pd
from docx import Document
from openpyxl import load_workbook
from PIL import Image,ImageOps,ImageDraw
ROOT=Path('/home/nikitha07/Projects/liver_graph'); OUT=ROOT/'manuscript_planning_2026-10-09_audit'; (OUT/'audit').mkdir(exist_ok=True)
files=sorted(p for p in ROOT.rglob('*') if p.is_file() and not any(x in p.parts for x in ['.git','.agents','.aws','.codex',OUT.name,'manuscript_planning_2026-10-09','__pycache__']))
records=[]; tables={}; textidx=[]; books={}
for p in files:
 rel=str(p.relative_to(ROOT)); b=p.read_bytes(); r={'path':rel,'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest(),'type':p.suffix,'status':'read_bytes'}
 try:
  if p.suffix.lower() in ['.csv','.tsv']:
   d=pd.read_csv(p,sep='\t' if p.suffix=='.tsv' else ',',low_memory=False); r.update(status='all_rows_parsed',rows=len(d),columns=list(d.columns)); tables[rel]={'rows':len(d),'columns':list(d.columns),'head':d.head(4).fillna('').to_dict('records')}
  elif p.suffix.lower() in ['.md','.txt','.py','.cypher','.html'] or p.name=='.gitignore':
   t=b.decode('utf-8-sig'); r.update(status='full_text_extracted',lines=len(t.splitlines())); heads=[{'line':i,'heading':l} for i,l in enumerate(t.splitlines(),1) if l.startswith('#')];textidx.append({'path':rel,'lines':r['lines'],'headings':heads})
  elif p.suffix.lower()=='.docx':
   doc=Document(p); txt='\n'.join([x.text for x in doc.paragraphs]+['\n'.join('\t'.join(c.text for c in row.cells) for row in table.rows) for table in doc.tables]); (OUT/'audit'/f'{p.stem}_extracted.txt').write_text(txt);r.update(status='paragraphs_and_tables_extracted',paragraphs=len(doc.paragraphs),tables=len(doc.tables))
  elif p.suffix.lower()=='.xlsx':
   wb=load_workbook(p,read_only=True,data_only=False); sheets=[]
   for ws in wb:
    vals=list(ws.values); sheets.append({'name':ws.title,'rows':len(vals),'columns':ws.max_column,'header':vals[0] if vals else [],'head':vals[1:5]})
    with (OUT/'audit'/f'{p.stem}__{ws.title}.csv').open('w') as f:csv.writer(f).writerows(vals)
   books[rel]=sheets;r.update(status='all_sheets_all_cells_extracted',sheets=[s['name'] for s in sheets]);wb.close()
  elif p.suffix.lower()=='.png':
   im=Image.open(p); im.load();r.update(status='decoded',width=im.width,height=im.height,dpi=im.info.get('dpi'))
 except Exception as e:r.update(status='failed',error=str(e))
 records.append(r)
(OUT/'audit/file_inventory.json').write_text(json.dumps(records,indent=2,default=str));(OUT/'audit/table_profiles.json').write_text(json.dumps(tables,indent=2,default=str));(OUT/'audit/workbook_profiles.json').write_text(json.dumps(books,indent=2,default=str));(OUT/'audit/text_section_index.json').write_text(json.dumps(textidx,indent=2))
manifest=pd.read_csv(ROOT/'Paper1_Package/00_START_HERE/06_FILE_MANIFEST.csv'); checks=[]
for _,row in manifest.iterrows():
 p=ROOT/'Paper1_Package'/row['path'];checks.append({'path':row['path'],'exists':p.exists(),'hash_matches':p.exists() and hashlib.sha256(p.read_bytes()).hexdigest()==row.sha256})
(OUT/'audit/manifest_check.json').write_text(json.dumps(checks,indent=2))
imgs=[p for p in files if 'Paper1_Package/01_FIGURES/' in str(p) and p.suffix=='.png']
for start in range(0,len(imgs),6):
 group=imgs[start:start+6];canvas=Image.new('RGB',(1600,1200),'white'); draw=ImageDraw.Draw(canvas)
 for i,p in enumerate(group):
  im=Image.open(p).convert('RGB');im.thumbnail((780,360));x=(i%2)*800;y=(i//2)*400;canvas.paste(im,(x+(800-im.width)//2,y+30));draw.text((x+10,y+8),p.name,fill='black')
 canvas.save(OUT/'audit'/f'figure_contact_sheet_{start//6+1}.png')
print('Files',len(records),'Package',sum(x['path'].startswith('Paper1_Package/') for x in records),'Types',dict(collections.Counter(x['type'] for x in records)))
print('Failed',[x for x in records if x['status']=='failed']);print('Manifest',len(checks),'mismatches',[x for x in checks if not x['hash_matches']])
print('Books',json.dumps(books,default=str))
print('Package figures',[{k:r[k] for k in ['path','width','height','dpi']} for r in records if r['path'].startswith('Paper1_Package/01_FIGURES/') and r['type']=='.png'])
